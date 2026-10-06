#!/usr/local/bin/python
"""Luiten CRM - webapp (Flask + SQLite), alleen standaard Python + Flask/openpyxl uit de FreeBSD ports.

Productie: draait achter Caddy (HTTPS) via het rc.d-script in deploy/, met LCRM_PUBLIC_URL=https://<domein>.
Lokaal testen: ./app.py  (http op poort 8000, alleen voor een vertrouwd netwerk)
"""
import csv, io, json, os, re, secrets, socket, sqlite3, sys, threading, time, zipfile, datetime as dt
from functools import lru_cache, wraps
from flask import Flask, abort, g, jsonify, redirect, request, send_file, send_from_directory, session
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash

APP_VERSION = '2026.10.1'
ENV = os.environ.get
BASE = os.path.dirname(os.path.abspath(__file__))
STATIC = os.path.join(BASE, 'static')
DATA = ENV('LCRM_DATA', os.path.join(BASE, 'data'))
UPLOADS = os.path.join(DATA, 'uploads')
DB_PATH = os.path.join(DATA, 'luiten-crm.db')
PORT = int(ENV('LCRM_PORT', '8000'))
PUBLIC_URL = ENV('LCRM_PUBLIC_URL', '').rstrip('/')
PUBLIC = PUBLIC_URL.startswith('https://')
HOST = ENV('LCRM_HOST', '127.0.0.1' if PUBLIC else '0.0.0.0')
SESSION_DAYS = int(ENV('LCRM_SESSION_DAYS', '7'))
MIN_PW, MAX_UPLOAD_MB = 10, 12
os.makedirs(UPLOADS, exist_ok=True)

STAGES = {'lev': ('lead', 'specs', 'prijs', 'proef', 'vast'), 'klant': ('lead', 'offerte', 'onderh', 'won', 'lost')}
ROLES = ('beheerder', 'inkoop', 'verkoop')
PRIOS = ('hot', 'warm', 'cold')
GROUPS = ('beef', 'lamb', 'poultry', 'game', 'pork', 'duck', 'other')
KINDS = ('photo', 'card', 'scan')
ID_RE = re.compile(r'^[A-Za-z0-9-]{8,40}$')
USER_RE = re.compile(r'^[a-z0-9._-]{2,32}$')
TIME_RE = re.compile(r'^([01]\d|2[0-3]):[0-5]\d$')
NOTE_FIELDS = {'hall': 20, 'stand': 40, 'contact': 120, 'contact_role': 120, 'contact_email': 160, 'contact_phone': 60,
               'text': 20000, 'price': 500}
DEFAULT_SETTINGS = {'event_name': 'SIAL Paris 2026', 'event_place': 'Paris Nord Villepinte', 'event_start': '2026-10-17', 'event_days': '5'}

SCHEMA = """
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL COLLATE NOCASE, name TEXT NOT NULL,
  initials TEXT NOT NULL, role TEXT NOT NULL, pw_hash TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1,
  must_change INTEGER NOT NULL DEFAULT 1, session_ver INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, last_seen TEXT);
CREATE TABLE IF NOT EXISTS relations(id INTEGER PRIMARY KEY, name TEXT NOT NULL, type TEXT NOT NULL DEFAULT 'lev',
  country TEXT NOT NULL DEFAULT '', owner_id INTEGER REFERENCES users(id), stage TEXT NOT NULL DEFAULT 'lead',
  value INTEGER NOT NULL DEFAULT 0, csb TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS notes(id TEXT PRIMARY KEY, relation_id INTEGER NOT NULL REFERENCES relations(id),
  user_id INTEGER NOT NULL REFERENCES users(id), date TEXT NOT NULL, time TEXT NOT NULL, hall TEXT NOT NULL DEFAULT '',
  stand TEXT NOT NULL DEFAULT '', contact TEXT NOT NULL DEFAULT '', contact_role TEXT NOT NULL DEFAULT '',
  contact_email TEXT NOT NULL DEFAULT '', contact_phone TEXT NOT NULL DEFAULT '', groups TEXT NOT NULL DEFAULT '',
  priority TEXT NOT NULL DEFAULT 'warm', text TEXT NOT NULL DEFAULT '', price TEXT NOT NULL DEFAULT '',
  created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS attachments(id TEXT PRIMARY KEY, note_id TEXT NOT NULL REFERENCES notes(id) ON DELETE CASCADE,
  kind TEXT NOT NULL, filename TEXT NOT NULL, created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS tasks(id TEXT PRIMARY KEY, note_id TEXT, relation_id INTEGER REFERENCES relations(id),
  text TEXT NOT NULL, owner_id INTEGER NOT NULL REFERENCES users(id), due TEXT NOT NULL DEFAULT '',
  done INTEGER NOT NULL DEFAULT 0, created_by INTEGER NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS meetings(id TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), date TEXT NOT NULL,
  time TEXT NOT NULL, company TEXT NOT NULL, country TEXT NOT NULL DEFAULT '', hall TEXT NOT NULL DEFAULT '',
  stand TEXT NOT NULL DEFAULT '', note_id TEXT, created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""
MIGRATE = {'users': {'must_change': 'INTEGER NOT NULL DEFAULT 0', 'session_ver': 'INTEGER NOT NULL DEFAULT 0'},
           'notes': {'contact_email': "TEXT NOT NULL DEFAULT ''", 'contact_phone': "TEXT NOT NULL DEFAULT ''"}}
INDEXES = """
CREATE INDEX IF NOT EXISTS ix_notes_user ON notes(user_id);
CREATE INDEX IF NOT EXISTS ix_notes_rel ON notes(relation_id);
CREATE INDEX IF NOT EXISTS ix_att_note ON attachments(note_id);
CREATE INDEX IF NOT EXISTS ix_att_file ON attachments(filename);
CREATE INDEX IF NOT EXISTS ix_tasks_owner ON tasks(owner_id);
CREATE INDEX IF NOT EXISTS ix_tasks_rel ON tasks(relation_id);
CREATE INDEX IF NOT EXISTS ix_meet_user ON meetings(user_id);
CREATE INDEX IF NOT EXISTS ix_rel_owner ON relations(owner_id);
CREATE INDEX IF NOT EXISTS ix_rel_name ON relations(name COLLATE NOCASE);
"""


# ---------- helpers ----------
def now():
    return dt.datetime.now().isoformat(timespec='seconds')


def clip(v, n=200):
    return str(v if v is not None else '').strip()[:n]


def err(msg, code=400, **kw):
    return jsonify(error=msg, **kw), code


def body():
    d = request.get_json(silent=True)
    return d if isinstance(d, dict) else {}


def valid_date(s):
    try:
        return dt.date.fromisoformat(str(s)[:10]).isoformat()
    except ValueError:
        return None


def initials(name):
    p = [x for x in re.split(r'[\s-]+', name) if x[:1].isalpha()]
    return ((p[0][0] + (p[-1][0] if len(p) > 1 else p[0][1:2])) if p else '??').upper()


def write_private(path, text):
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, 'w', encoding='utf-8') as f:
        f.write(text)


def secret_key():
    p = os.path.join(DATA, 'secret.key')
    if not os.path.exists(p):
        write_private(p, secrets.token_hex(32))
    with open(p) as f:
        return f.read().strip()


def pw_problem(pw, *forbidden):
    pw = str(pw or '')
    if len(pw) < MIN_PW:
        return 'Wachtwoord moet minstens %d tekens zijn' % MIN_PW
    if len(pw) > 128:
        return 'Wachtwoord is te lang'
    if len(set(pw)) < 4 or pw.lower() in {str(f).lower() for f in forbidden if f}:
        return 'Kies een minder voorspelbaar wachtwoord'
    return None


# ---------- app + security ----------
app = Flask(__name__, static_folder=STATIC, static_url_path='/static')
app.secret_key = secret_key()
app.config.update(MAX_CONTENT_LENGTH=MAX_UPLOAD_MB * 1024 * 1024, PERMANENT_SESSION_LIFETIME=dt.timedelta(days=SESSION_DAYS),
                  SESSION_COOKIE_NAME='lcrm', SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                  SESSION_COOKIE_SECURE=PUBLIC, JSON_AS_ASCII=False)
if PUBLIC:
    from werkzeug.middleware.proxy_fix import ProxyFix
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

CSP = ("default-src 'self'; img-src 'self' blob: data:; style-src 'self' 'unsafe-inline'; script-src 'self'; "
       "connect-src 'self'; font-src 'self'; object-src 'none'; base-uri 'none'; form-action 'self'; frame-ancestors 'none'")
HTTP_MSG = {400: 'Ongeldig verzoek', 401: 'Niet ingelogd', 403: 'Geweigerd', 404: 'Niet gevonden', 405: 'Niet toegestaan',
            413: 'Bestand is te groot (max %d MB)' % MAX_UPLOAD_MB, 429: 'Te veel verzoeken'}


@app.before_request
def csrf_guard():
    """Wijzigingen alleen vanaf de eigen app: eigen header (niet cross-site te zetten) + Origin-controle."""
    if request.method in ('GET', 'HEAD', 'OPTIONS'):
        return None
    origin = request.headers.get('Origin')
    if request.headers.get('X-LCRM') != '1' or (origin and origin.rstrip('/') != request.host_url.rstrip('/')):
        return err('Geweigerd', 403)
    return None


@app.after_request
def security_headers(r):
    h = r.headers
    h['Content-Security-Policy'] = CSP
    h['X-Content-Type-Options'] = 'nosniff'
    h['X-Frame-Options'] = 'DENY'
    h['Referrer-Policy'] = 'same-origin'
    h['Permissions-Policy'] = 'camera=(self), microphone=(), geolocation=(), payment=()'
    h['Cross-Origin-Opener-Policy'] = 'same-origin'
    h['Cross-Origin-Resource-Policy'] = 'same-origin'
    if PUBLIC:
        h['Strict-Transport-Security'] = 'max-age=31536000; includeSubDomains'
    if request.path.startswith('/api/') and 'Cache-Control' not in h:
        h['Cache-Control'] = 'no-store'
    return r


@app.errorhandler(Exception)
def on_error(e):
    code = e.code if isinstance(e, HTTPException) else 500
    if code == 500:
        app.logger.exception(e)
    if request.path.startswith('/api/') or code == 500:
        return err(HTTP_MSG.get(code, 'Serverfout'), code)
    return e


# ---------- database ----------
VERSION = [secrets.randbits(31)]  # verandert bij elke wijziging -> ETag voor /api/data


def db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH, timeout=10)
        g.db.row_factory = sqlite3.Row
        g.db.execute('PRAGMA foreign_keys=ON')
    return g.db


@app.teardown_appcontext
def close_db(_):
    c = g.pop('db', None)
    if c:
        c.close()


def q(sql, a=()):
    return [dict(r) for r in db().execute(sql, a)]


def q1(sql, a=()):
    r = db().execute(sql, a).fetchone()
    return dict(r) if r else None


def ex(sql, a=()):
    return db().execute(sql, a)


def commit():
    db().commit()
    VERSION[0] += 1


def settings():
    return {r['key']: r['value'] for r in q('SELECT key, value FROM settings')}


def init_db():
    con = sqlite3.connect(DB_PATH)
    con.executescript(SCHEMA)
    for table, cols in MIGRATE.items():
        have = {r[1] for r in con.execute('PRAGMA table_info(%s)' % table)}
        for c, decl in cols.items():
            if c not in have:
                con.execute('ALTER TABLE %s ADD COLUMN %s %s' % (table, c, decl))
    con.executescript(INDEXES)
    con.executemany('INSERT OR IGNORE INTO settings(key, value) VALUES(?, ?)', DEFAULT_SETTINGS.items())
    first = None
    if not con.execute('SELECT 1 FROM users LIMIT 1').fetchone():
        first = secrets.token_urlsafe(9)
        con.execute('INSERT INTO users(username, name, initials, role, pw_hash, must_change, created_at) VALUES(?,?,?,?,?,1,?)',
                    ('beheerder', 'Beheerder', 'BE', 'beheerder', generate_password_hash(first), now()))
        write_private(os.path.join(DATA, 'EERSTE-WACHTWOORD.txt'),
                      'Gebruikersnaam: beheerder\nWachtwoord: %s\n\nBij de eerste keer inloggen kies je een eigen wachtwoord.\n'
                      'Verwijder daarna dit bestand.\n' % first)
    con.commit()  # journal_mode kan niet binnen een open transactie wijzigen (anders faalt elke herstart)
    con.execute('PRAGMA journal_mode=WAL')
    con.close()
    return first


# ---------- authenticatie ----------
FAILS, FAIL_LOCK, DUMMY_HASH = {}, threading.Lock(), generate_password_hash(secrets.token_hex(16))
LIMITS = {'ip': 20, 'user': 8}  # mislukte pogingen per 15 minuten


def throttled(keys, add=False):
    t = time.time()
    with FAIL_LOCK:
        if len(FAILS) > 10000:
            FAILS.clear()
        for k in keys:
            hits = [x for x in FAILS.get(k, ()) if t - x < 900] + ([t] if add else [])
            FAILS[k] = hits
            if not add and len(hits) >= LIMITS[k[0]]:
                return True
    return False


def me():
    if 'me' not in g:
        uid = session.get('uid')
        u = q1('SELECT id, username, name, initials, role, must_change, session_ver FROM users WHERE id=? AND active=1', (uid,)) if uid else None
        g.me = u if u and u['session_ver'] == session.get('sv') else None
    return g.me


def is_admin(u):
    return u['role'] == 'beheerder'


def can(u, owner_id):
    return is_admin(u) or owner_id == u['id']


def auth(admin=False, allow_change=False):
    def deco(fn):
        @wraps(fn)
        def wrap(*a, **kw):
            u = me()
            if not u:
                return err('Niet ingelogd', 401)
            if u['must_change'] and not allow_change:
                return err('Kies eerst een nieuw wachtwoord', 403, must_change=True)
            if admin and not is_admin(u):
                return err('Alleen voor beheerder', 403)
            return fn(u, *a, **kw)
        return wrap
    return deco


@app.post('/api/login')
def login():
    d = body()
    name, pw = clip(d.get('username'), 64).lower(), str(d.get('password') or '')[:128]
    keys = [('ip', request.remote_addr or '?'), ('user', name)]
    if throttled(keys):
        return err('Te veel mislukte pogingen. Probeer het over 15 minuten opnieuw.', 429)
    u = q1('SELECT * FROM users WHERE username=? AND active=1', (name,))
    ok = check_password_hash(u['pw_hash'] if u else DUMMY_HASH, pw)
    if not (u and ok):
        throttled(keys, add=True)
        return err('Onjuiste gebruikersnaam of wachtwoord', 401)
    session.clear()
    session.permanent = True
    session.update(uid=u['id'], sv=u['session_ver'])
    ex('UPDATE users SET last_seen=? WHERE id=?', (now(), u['id']))
    db().commit()
    return jsonify(ok=True, role=u['role'], must_change=bool(u['must_change']))


@app.post('/api/logout')
def logout():
    session.clear()
    return jsonify(ok=True)


@app.post('/api/password')
@auth(allow_change=True)
def change_password(u):
    d = body()
    old, new = str(d.get('old') or '')[:128], str(d.get('new') or '')
    keys = [('ip', request.remote_addr or '?'), ('user', u['username'])]
    if throttled(keys):
        return err('Te veel mislukte pogingen. Probeer het over 15 minuten opnieuw.', 429)
    if not check_password_hash(q1('SELECT pw_hash FROM users WHERE id=?', (u['id'],))['pw_hash'], old):
        throttled(keys, add=True)
        return err('Huidig wachtwoord klopt niet')
    problem = pw_problem(new, old, u['username'], 'wachtwoord', 'password', 'luitenfood', 'welkom123')
    if problem:
        return err(problem)
    ex('UPDATE users SET pw_hash=?, must_change=0, session_ver=session_ver+1 WHERE id=?', (generate_password_hash(new), u['id']))
    commit()
    session['sv'] = u['session_ver'] + 1  # andere apparaten worden uitgelogd, dit apparaat blijft ingelogd
    return jsonify(ok=True)


# ---------- pagina's ----------
def page(name):
    return send_from_directory(STATIC, name, max_age=0)


def gate(target):
    u = me()
    if not u:
        return redirect('/login')
    if u['must_change']:
        return redirect('/login?change=1')
    if target == 'admin' and not is_admin(u):
        return redirect('/app')
    return page(target + '.html')


@app.get('/')
def home():
    u = me()
    return redirect('/login' if not u else '/admin' if is_admin(u) else '/app')


app.add_url_rule('/login', 'login_page', lambda: page('login.html'))
app.add_url_rule('/app', 'app_page', lambda: gate('app'))
app.add_url_rule('/admin', 'admin_page', lambda: gate('admin'))


# ---------- health (voor de statuspagina, zonder login) ----------
@app.get('/health')
def health():
    checks = {}
    try:
        con = sqlite3.connect('file:%s?mode=ro' % DB_PATH, uri=True, timeout=3)
        try:
            con.execute('SELECT 1 FROM users LIMIT 1').fetchone()
        finally:
            con.close()
        checks['database'] = 'ok'
    except sqlite3.Error:
        checks['database'] = 'error'
    checks['storage'] = 'ok' if os.access(UPLOADS, os.W_OK) else 'error'
    ok = all(v == 'ok' for v in checks.values())
    r = jsonify(status='ok' if ok else 'error', version=APP_VERSION, **checks,
                time=dt.datetime.now().astimezone().isoformat(timespec='seconds'))
    r.status_code = 200 if ok else 503
    r.headers['Cache-Control'] = 'no-store'
    return r


# ---------- data ----------
@lru_cache(maxsize=1)
def lan_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(('10.255.255.255', 1))
        return s.getsockname()[0]
    except OSError:
        return '127.0.0.1'
    finally:
        s.close()


def app_url():
    return PUBLIC_URL or 'http://%s:%d' % (lan_ip(), PORT)


NOTES_SQL = ('SELECT n.*, r.name AS company, r.type, r.country, r.stage FROM notes n JOIN relations r ON r.id=n.relation_id')


def visible_data(u):
    adm, uid = is_admin(u), u['id']
    nw, na = ('', ()) if adm else (' WHERE n.user_id=?', (uid,))
    notes = q(NOTES_SQL + nw + ' ORDER BY n.date DESC, n.time DESC', na)
    atts = {}
    for a in q('SELECT a.id, a.note_id, a.kind, a.filename FROM attachments a JOIN notes n ON n.id=a.note_id' + nw + ' ORDER BY a.created_at', na):
        atts.setdefault(a['note_id'], []).append({'id': a['id'], 'kind': a['kind'], 'url': '/media/' + a['filename']})
    for n in notes:
        n['groups'] = n['groups'].split(',') if n['groups'] else []
        n['attachments'] = atts.get(n['id'], [])
    tasks = q('SELECT t.*, r.name AS company FROM tasks t LEFT JOIN relations r ON r.id=t.relation_id'
              + ('' if adm else ' WHERE t.owner_id=? OR t.created_by=?') + ' ORDER BY t.done, t.due', () if adm else (uid, uid))
    for t in tasks:
        t['done'] = bool(t['done'])
    meetings = q('SELECT * FROM meetings' + ('' if adm else ' WHERE user_id=?') + ' ORDER BY date, time', () if adm else (uid,))
    relations = q('SELECT r.*, COUNT(n.id) AS note_count, MAX(n.date) AS last_note FROM relations r LEFT JOIN notes n ON n.relation_id=r.id'
                  + ('' if adm else ' AND n.user_id=? WHERE r.owner_id=? OR n.id IS NOT NULL')
                  + ' GROUP BY r.id ORDER BY r.name COLLATE NOCASE', () if adm else (uid, uid))
    return {'notes': notes, 'tasks': tasks, 'meetings': meetings, 'relations': relations}


@app.get('/api/me')
@auth(allow_change=True)
def api_me(u):
    ex('UPDATE users SET last_seen=? WHERE id=?', (now(), u['id']))
    db().commit()
    users = q('SELECT id, name, initials, role FROM users WHERE active=1 ORDER BY name COLLATE NOCASE')
    user = {k: u[k] for k in ('id', 'username', 'name', 'initials', 'role')}
    return jsonify(user={**user, 'must_change': bool(u['must_change'])}, users=users, settings=settings(), url=app_url())


@app.get('/api/data')
@auth()
def api_data(u):
    tag = '"%x-%d"' % (VERSION[0], u['id'])
    if request.headers.get('If-None-Match') == tag:
        return '', 304
    r = jsonify(visible_data(u))
    r.headers['ETag'] = tag
    return r


def relation_for(name, typ, country, uid):
    r = q1('SELECT id, owner_id FROM relations WHERE name=? COLLATE NOCASE', (name,))
    if r:
        return r
    t = now()
    cur = ex('INSERT INTO relations(name, type, country, owner_id, created_at, updated_at) VALUES(?,?,?,?,?,?)', (name, typ, country, uid, t, t))
    return {'id': cur.lastrowid, 'owner_id': uid}


@app.post('/api/notes')
@auth()
def save_note(u):
    d = body()
    nid, company = clip(d.get('id'), 40), clip(d.get('company'), 160)
    if not ID_RE.match(nid):
        return err('Ongeldig id')
    if not company:
        return err('Bedrijf is verplicht')
    typ = d.get('type') if d.get('type') in STAGES else 'lev'
    stage = d.get('stage') if d.get('stage') in STAGES[typ] else 'lead'
    prio = d.get('priority') if d.get('priority') in PRIOS else 'warm'
    raw = d.get('groups') if isinstance(d.get('groups'), list) else []
    groups = ','.join(dict.fromkeys(x for x in raw if x in GROUPS)) or 'other'
    old = q1('SELECT user_id FROM notes WHERE id=?', (nid,))
    if old and not can(u, old['user_id']):
        return err('Geen rechten op deze notitie', 403)
    country = clip(d.get('country'), 80)
    rel = relation_for(company, typ, country, u['id'])
    if can(u, rel['owner_id']):
        ex("UPDATE relations SET type=?, stage=?, country=COALESCE(NULLIF(?, ''), country), updated_at=? WHERE id=?",
           (typ, stage, country, now(), rel['id']))
    f = {k: clip(d.get(k), n) for k, n in NOTE_FIELDS.items()}
    t = now()
    if old:
        ex('UPDATE notes SET relation_id=?, groups=?, priority=?, updated_at=?, %s WHERE id=?' % ', '.join(k + '=?' for k in f),
           (rel['id'], groups, prio, t, *f.values(), nid))
    else:
        date = valid_date(d.get('date')) or dt.date.today().isoformat()
        tm = d.get('time') if TIME_RE.match(str(d.get('time') or '')) else dt.datetime.now().strftime('%H:%M')
        cols = ('id', 'relation_id', 'user_id', 'date', 'time', 'groups', 'priority', 'created_at', 'updated_at', *f)
        ex('INSERT INTO notes(%s) VALUES(%s)' % (','.join(cols), ','.join('?' * len(cols))),
           (nid, rel['id'], u['id'], date, tm, groups, prio, t, t, *f.values()))
    if d.get('meeting_id'):
        ex('UPDATE meetings SET note_id=? WHERE id=? AND (user_id=? OR ?)', (nid, clip(d['meeting_id'], 40), u['id'], int(is_admin(u))))
    if isinstance(d.get('task'), dict) and clip(d['task'].get('text')):
        upsert_task(d['task'], u, note_id=nid, relation_id=rel['id'])
    commit()
    return jsonify(ok=True, id=nid, relation_id=rel['id'])


def remove_files(names):
    for fn in names:
        try:
            os.remove(os.path.join(UPLOADS, fn))
        except OSError:
            pass


@app.delete('/api/notes/<nid>')
@auth()
def delete_note(u, nid):
    n = q1('SELECT user_id FROM notes WHERE id=?', (nid,))
    if not n:
        return jsonify(ok=True)
    if not can(u, n['user_id']):
        return err('Geen rechten', 403)
    remove_files(r['filename'] for r in q('SELECT filename FROM attachments WHERE note_id=?', (nid,)))
    ex('DELETE FROM attachments WHERE note_id=?', (nid,))
    ex('UPDATE tasks SET note_id=NULL WHERE note_id=?', (nid,))
    ex('UPDATE meetings SET note_id=NULL WHERE note_id=?', (nid,))
    ex('DELETE FROM notes WHERE id=?', (nid,))
    commit()
    return jsonify(ok=True)


def image_ext(head):
    if head.startswith(b'\xff\xd8\xff'):
        return '.jpg'
    if head.startswith(b'\x89PNG\r\n\x1a\n'):
        return '.png'
    if head[:4] == b'RIFF' and head[8:12] == b'WEBP':
        return '.webp'
    return None


@app.post('/api/notes/<nid>/attachments')
@auth()
def upload(u, nid):
    n = q1('SELECT user_id FROM notes WHERE id=?', (nid,))
    if not n:
        return err('Notitie niet gevonden', 404)
    if not can(u, n['user_id']):
        return err('Geen rechten', 403)
    f, aid = request.files.get('file'), clip(request.form.get('id'), 40)
    if not f or not ID_RE.match(aid) or not ID_RE.match(nid):
        return err('Ongeldige upload')
    if q1('SELECT 1 FROM attachments WHERE id=?', (aid,)):
        return jsonify(ok=True, id=aid)
    data = f.read()
    ext = image_ext(data[:12])
    if not ext:
        return err('Alleen JPG-, PNG- of WebP-foto’s')
    kind = request.form.get('kind') if request.form.get('kind') in KINDS else 'photo'
    fn = '%s_%s%s' % (nid, aid, ext)
    with open(os.path.join(UPLOADS, fn), 'wb') as out:
        out.write(data)
    ex('INSERT INTO attachments(id, note_id, kind, filename, created_at) VALUES(?,?,?,?,?)', (aid, nid, kind, fn, now()))
    commit()
    return jsonify(ok=True, id=aid, url='/media/' + fn)


@app.delete('/api/attachments/<aid>')
@auth()
def delete_attachment(u, aid):
    a = q1('SELECT a.filename, n.user_id FROM attachments a JOIN notes n ON n.id=a.note_id WHERE a.id=?', (aid,))
    if not a:
        return jsonify(ok=True)
    if not can(u, a['user_id']):
        return err('Geen rechten', 403)
    remove_files([a['filename']])
    ex('DELETE FROM attachments WHERE id=?', (aid,))
    commit()
    return jsonify(ok=True)


@app.get('/media/<fn>')
def media(fn):
    u = me()
    if not u:
        abort(401)
    a = q1('SELECT n.user_id FROM attachments a JOIN notes n ON n.id=a.note_id WHERE a.filename=?', (fn,))
    if not a or not can(u, a['user_id']):
        abort(404)
    r = send_from_directory(UPLOADS, fn, max_age=86400)
    r.headers['Cache-Control'] = 'private, max-age=86400'
    return r


def upsert_task(d, u, note_id=None, relation_id=None):
    tid = clip(d.get('id'), 40) or secrets.token_hex(8)
    if not ID_RE.match(tid):
        return None
    try:
        owner = int(d.get('owner_id') or u['id'])
    except (TypeError, ValueError):
        owner = u['id']
    if not q1('SELECT 1 FROM users WHERE id=? AND active=1', (owner,)):
        owner = u['id']
    due = valid_date(d.get('due')) or ''
    old = q1('SELECT owner_id, created_by FROM tasks WHERE id=?', (tid,))
    t = now()
    if old:
        if not (is_admin(u) or u['id'] in (old['owner_id'], old['created_by'])):
            return None
        ex('UPDATE tasks SET text=?, owner_id=?, due=?, done=?, updated_at=? WHERE id=?',
           (clip(d.get('text'), 500), owner, due, int(bool(d.get('done'))), t, tid))
        return tid
    if relation_id is None and d.get('relation_id'):
        r = q1('SELECT id, owner_id FROM relations WHERE id=?', (d['relation_id'],))
        relation_id = r['id'] if r and can(u, r['owner_id']) else None
    ex('INSERT INTO tasks(id, note_id, relation_id, text, owner_id, due, done, created_by, created_at, updated_at) VALUES(?,?,?,?,?,?,0,?,?,?)',
       (tid, note_id, relation_id, clip(d.get('text'), 500), owner, due, u['id'], t, t))
    return tid


@app.post('/api/tasks')
@auth()
def save_task(u):
    d = body()
    if not clip(d.get('text')):
        return err('Tekst is verplicht')
    tid = upsert_task(d, u)
    if not tid:
        return err('Geen rechten', 403)
    commit()
    return jsonify(ok=True, id=tid)


@app.patch('/api/tasks/<tid>')
@auth()
def toggle_task(u, tid):
    t = q1('SELECT owner_id, created_by FROM tasks WHERE id=?', (tid,))
    if not t:
        return err('Actie niet gevonden', 404)
    if not (is_admin(u) or u['id'] in (t['owner_id'], t['created_by'])):
        return err('Geen rechten', 403)
    ex('UPDATE tasks SET done=?, updated_at=? WHERE id=?', (int(bool(body().get('done'))), now(), tid))
    commit()
    return jsonify(ok=True)


@app.delete('/api/tasks/<tid>')
@auth(admin=True)
def delete_task(u, tid):
    ex('DELETE FROM tasks WHERE id=?', (tid,))
    commit()
    return jsonify(ok=True)


@app.post('/api/meetings')
@auth()
def save_meeting(u):
    d = body()
    mid, date, company = clip(d.get('id'), 40), valid_date(d.get('date')), clip(d.get('company'), 160)
    if not ID_RE.match(mid) or not date or not company:
        return err('Bedrijf en datum zijn verplicht')
    tm = d.get('time') if TIME_RE.match(str(d.get('time') or '')) else '09:00'
    vals = (date, tm, company, clip(d.get('country'), 80), clip(d.get('hall'), 20), clip(d.get('stand'), 40))
    old = q1('SELECT user_id FROM meetings WHERE id=?', (mid,))
    if old:
        if not can(u, old['user_id']):
            return err('Geen rechten', 403)
        ex('UPDATE meetings SET date=?, time=?, company=?, country=?, hall=?, stand=? WHERE id=?', (*vals, mid))
    else:
        ex('INSERT INTO meetings(id, user_id, date, time, company, country, hall, stand, created_at) VALUES(?,?,?,?,?,?,?,?,?)',
           (mid, u['id'], *vals, now()))
    commit()
    return jsonify(ok=True)


@app.delete('/api/meetings/<mid>')
@auth()
def delete_meeting(u, mid):
    m = q1('SELECT user_id FROM meetings WHERE id=?', (mid,))
    if m and can(u, m['user_id']):
        ex('DELETE FROM meetings WHERE id=?', (mid,))
        commit()
    return jsonify(ok=True)


# ---------- beheer ----------
@app.get('/api/admin/users')
@auth(admin=True)
def admin_users(u):
    us = q('SELECT id, username, name, initials, role, active, created_at, last_seen FROM users ORDER BY active DESC, name COLLATE NOCASE')
    for x in us:
        x['active'] = bool(x['active'])
    return jsonify(us)


@app.post('/api/admin/users')
@auth(admin=True)
def admin_create_user(u):
    d = body()
    name, username, role, pw = clip(d.get('name'), 80), clip(d.get('username'), 32).lower(), d.get('role'), str(d.get('password') or '')
    if not name or not USER_RE.match(username):
        return err('Vul een naam in en een gebruikersnaam van 2-32 tekens (a-z, 0-9, . _ -)')
    if role not in ROLES:
        return err('Kies een rol')
    problem = pw_problem(pw, username)
    if problem:
        return err(problem)
    if q1('SELECT 1 FROM users WHERE username=?', (username,)):
        return err('Gebruikersnaam bestaat al')
    ex('INSERT INTO users(username, name, initials, role, pw_hash, must_change, created_at) VALUES(?,?,?,?,?,1,?)',
       (username, name, initials(name), role, generate_password_hash(pw), now()))
    commit()
    return jsonify(ok=True)


@app.patch('/api/admin/users/<int:uid>')
@auth(admin=True)
def admin_update_user(u, uid):
    d = body()
    if not q1('SELECT 1 FROM users WHERE id=?', (uid,)):
        return err('Gebruiker niet gevonden', 404)
    if uid == u['id'] and ({'role', 'active', 'password'} & d.keys()):
        return err('Wijzig je eigen account via de app (tandwiel)')
    sets, args = [], []
    if 'role' in d:
        if d['role'] not in ROLES:
            return err('Ongeldige rol')
        sets.append('role=?')
        args.append(d['role'])
    if 'active' in d:
        sets.append('active=?')
        args.append(int(bool(d['active'])))
    if clip(d.get('name')):
        n = clip(d['name'], 80)
        sets += ['name=?', 'initials=?']
        args += [n, initials(n)]
    if d.get('password'):
        problem = pw_problem(d['password'])
        if problem:
            return err(problem)
        sets += ['pw_hash=?', 'must_change=1']
        args.append(generate_password_hash(str(d['password'])))
    if not sets:
        return jsonify(ok=True)
    if {'role', 'active', 'password'} & d.keys():
        sets.append('session_ver=session_ver+1')  # bestaande sessies van deze gebruiker vervallen
    ex('UPDATE users SET %s WHERE id=?' % ', '.join(sets), (*args, uid))
    commit()
    return jsonify(ok=True)


def relation_values(d, r=None):
    r = r or {'name': '', 'type': 'lev', 'country': '', 'owner_id': None, 'stage': 'lead', 'value': 0, 'csb': ''}
    typ = d.get('type') if d.get('type') in STAGES else r['type']
    stage = d.get('stage', r['stage'])
    try:
        value = max(0, int(float(d.get('value', r['value']) or 0)))
    except (TypeError, ValueError):
        value = r['value']
    owner = d.get('owner_id', r['owner_id'])
    if not q1('SELECT 1 FROM users WHERE id=?', (owner,)):
        owner = r['owner_id']
    return (clip(d.get('name')) and clip(d['name'], 160)) or r['name'], typ, clip(d.get('country', r['country']), 80), owner, \
        stage if stage in STAGES[typ] else 'lead', value, clip(d.get('csb', r['csb']), 40)


@app.post('/api/relations')
@auth(admin=True)
def admin_create_relation(u):
    d = body()
    vals = relation_values({**d, 'owner_id': d.get('owner_id') or u['id']})
    if not vals[0]:
        return err('Naam is verplicht')
    if q1('SELECT 1 FROM relations WHERE name=? COLLATE NOCASE', (vals[0],)):
        return err('Deze relatie bestaat al')
    t = now()
    cur = ex('INSERT INTO relations(name, type, country, owner_id, stage, value, csb, created_at, updated_at) VALUES(?,?,?,?,?,?,?,?,?)',
             (*vals, t, t))
    commit()
    return jsonify(ok=True, id=cur.lastrowid)


@app.patch('/api/relations/<int:rid>')
@auth(admin=True)
def admin_update_relation(u, rid):
    r = q1('SELECT * FROM relations WHERE id=?', (rid,))
    if not r:
        return err('Relatie niet gevonden', 404)
    ex('UPDATE relations SET name=?, type=?, country=?, owner_id=?, stage=?, value=?, csb=?, updated_at=? WHERE id=?',
       (*relation_values(body(), r), now(), rid))
    commit()
    return jsonify(ok=True)


@app.post('/api/relations/<int:rid>/merge')
@auth(admin=True)
def admin_merge_relation(u, rid):
    try:
        into = int(body().get('into') or 0)
    except (TypeError, ValueError):
        into = 0
    if into == rid or not q1('SELECT 1 FROM relations WHERE id=?', (into,)):
        return err('Kies een andere relatie om mee samen te voegen')
    ex('UPDATE notes SET relation_id=? WHERE relation_id=?', (into, rid))
    ex('UPDATE tasks SET relation_id=? WHERE relation_id=?', (into, rid))
    ex('DELETE FROM relations WHERE id=?', (rid,))
    commit()
    return jsonify(ok=True)


@app.post('/api/admin/settings')
@auth(admin=True)
def admin_settings(u):
    d = body()
    for k in DEFAULT_SETTINGS:
        if k not in d:
            continue
        v = clip(d[k], 120)
        if k == 'event_start' and not valid_date(v):
            return err('Startdatum moet als JJJJ-MM-DD')
        if k == 'event_days' and not (v.isdigit() and 1 <= int(v) <= 14):
            return err('Aantal dagen tussen 1 en 14')
        ex('INSERT OR REPLACE INTO settings(key, value) VALUES(?, ?)', (k, v))
    commit()
    return jsonify(ok=True)


# ---------- export ----------
GROUP_NL = {'beef': 'Rund', 'lamb': 'Lam', 'poultry': 'Gevogelte', 'game': 'Wild', 'pork': 'Ibérico / varken', 'duck': 'Eend & gans', 'other': 'Overig'}
PRIO_NL = {'hot': 'Hot lead', 'warm': 'Warm', 'cold': 'Koud'}
STAGE_NL = {'lead': 'Lead', 'specs': 'Monsters & specs', 'prijs': 'Prijsonderhandeling', 'proef': 'Proeforder', 'vast': 'Vaste leverancier',
            'offerte': 'Offerte', 'onderh': 'Onderhandeling', 'won': 'Gewonnen', 'lost': 'Verloren'}
KIND_DIR = {'card': 'visitekaartjes', 'photo': 'fotos', 'scan': 'gescande-notities'}
KIND_NL = {'card': 'Visitekaartje', 'photo': 'Foto', 'scan': 'Gescande notitie'}
NOTE_COLS = [('date', 'Datum'), ('time', 'Tijd'), ('company', 'Bedrijf'), ('type', 'Type'), ('country', 'Land'), ('hall', 'Hal'),
             ('stand', 'Stand'), ('contact', 'Contactpersoon'), ('contact_role', 'Functie'), ('contact_email', 'E-mail'),
             ('contact_phone', 'Telefoon'), ('groups', 'Productgroepen'), ('priority', 'Prioriteit'), ('stage', 'Fase'),
             ('text', 'Notitie'), ('price', 'Prijs & volume'), ('user', 'Vastgelegd door'), ('attachments', 'Bijlagen')]
TASK_COLS = [('company', 'Relatie'), ('text', 'Actie'), ('owner', 'Eigenaar'), ('due', 'Deadline'), ('done', 'Klaar')]
REL_COLS = [('name', 'Relatie'), ('type', 'Type'), ('country', 'Land'), ('stage', 'Fase'), ('owner', 'Eigenaar'),
            ('value', 'Verwachte waarde (EUR/jaar)'), ('csb', 'CSB-nummer'), ('note_count', 'Notities'), ('last_note', 'Laatste notitie')]
MEET_COLS = [('date', 'Datum'), ('time', 'Tijd'), ('company', 'Bedrijf'), ('country', 'Land'), ('hall', 'Hal'), ('stand', 'Stand'),
             ('user', 'Medewerker'), ('logged', 'Notitie vastgelegd')]
IDX_COLS = [('relation', 'Relatie'), ('type', 'Type'), ('country', 'Land'), ('csb', 'CSB-nummer'), ('contact', 'Contactpersoon'),
            ('role', 'Functie'), ('email', 'E-mail'), ('phone', 'Telefoon'), ('kind', 'Soort'), ('date', 'Datum'), ('time', 'Tijd'),
            ('user', 'Vastgelegd door'), ('file', 'Bestand')]
NUMLIKE = re.compile(r'^[+-][\d\s().,/-]*$')


def safe_cell(v):
    """Voorkomt formule-injectie in Excel/CSV (=, +, -, @ aan het begin), maar laat telefoonnummers en getallen intact."""
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@', '\t', '\r') and not NUMLIKE.match(v):
        return "'" + v
    return v


def type_nl(t):
    return 'Klant' if t == 'klant' else 'Leverancier'


def export_tables(u, date=None):
    data = visible_data(u)
    names = {r['id']: r['name'] for r in q('SELECT id, name FROM users')}
    notes = [n for n in data['notes'] if not date or n['date'] == date]
    for n in notes:
        n.update(user=names.get(n['user_id'], ''), groups=', '.join(GROUP_NL.get(x, x) for x in n['groups']), type=type_nl(n['type']),
                 priority=PRIO_NL.get(n['priority'], n['priority']), stage=STAGE_NL.get(n['stage'], n['stage']), attachments=len(n['attachments']))
    for t in data['tasks']:
        t.update(owner=names.get(t['owner_id'], ''), done='ja' if t['done'] else 'nee')
    for r in data['relations']:
        r.update(owner=names.get(r['owner_id'], ''), type=type_nl(r['type']), stage=STAGE_NL.get(r['stage'], r['stage']))
    for m in data['meetings']:
        m.update(user=names.get(m['user_id'], ''), logged='ja' if m['note_id'] else 'nee')
    return {'Notities': (NOTE_COLS, notes), 'Acties': (TASK_COLS, data['tasks']),
            'Relaties': (REL_COLS, data['relations']), 'Afspraken': (MEET_COLS, data['meetings'])}


def to_csv(cols, items):
    buf = io.StringIO()
    buf.write('\ufeff')
    w = csv.writer(buf, delimiter=';')
    w.writerow([c[1] for c in cols])
    w.writerows([safe_cell(it.get(c[0], '')) for c in cols] for it in items)
    return buf.getvalue().encode('utf-8')


def to_xlsx(sheets, link_col=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill, link_font = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='022D4E'), Font(color='0563C1', underline='single')
    for title, (cols, items) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append([c[1] for c in cols])
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for it in items:
            ws.append([safe_cell(it.get(c[0], '')) for c in cols])
            if link_col:
                cell = ws.cell(ws.max_row, len(cols))
                cell.hyperlink, cell.font = it[link_col], link_font
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 60 if c[0] in ('text', 'file') else 18
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def stamp():
    return dt.datetime.now().strftime('%Y%m%d-%H%M')


def download(data, mime, name):
    return send_file(io.BytesIO(data), mimetype=mime, as_attachment=True, download_name=name)


@app.get('/api/export/csv')
@auth()
def export_csv(u):
    what = {'notes': 'Notities', 'tasks': 'Acties', 'relations': 'Relaties', 'meetings': 'Afspraken'}.get(request.args.get('what'), 'Notities')
    date = valid_date(request.args.get('date')) if request.args.get('date') else None
    cols, items = export_tables(u, date)[what]
    return download(to_csv(cols, items), 'text/csv', 'luiten-crm-%s%s-%s.csv' % (what.lower(), '-' + date if date else '', stamp()))


@app.get('/api/export/xlsx')
@auth()
def export_xlsx(u):
    date = valid_date(request.args.get('date')) if request.args.get('date') else None
    return download(to_xlsx(export_tables(u, date)), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    'luiten-crm-%s.xlsx' % stamp())


@app.get('/api/export/json')
@auth()
def export_json(u):
    payload = {'exported_at': now(), 'exported_by': u['name'], 'settings': settings(), **visible_data(u)}
    return download(json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8'), 'application/json', 'luiten-crm-%s.json' % stamp())


def safe_name(s, maxlen=80):
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '', str(s or ''))
    s = re.sub(r'\s+', ' ', s).strip().rstrip('. ')
    if re.fullmatch(r'(CON|PRN|AUX|NUL|COM\d|LPT\d)', s.upper()):
        s = '_' + s
    return s[:maxlen].rstrip('. ') or 'Onbekend'


def note_txt(n, names, tasks, files):
    contact = ', '.join(x for x in (n['contact'], n['contact_role'], n['contact_email'], n['contact_phone']) if x)
    loc = ' · '.join(x for x in (n['hall'] and 'Hal ' + n['hall'], n['stand'] and 'Stand ' + n['stand']) if x)
    L = ['=' * 64, '%s %s  ·  %s' % (n['date'], n['time'], names.get(n['user_id'], '')), '=' * 64]
    L += ['Contact: ' + contact] if contact else []
    L += [loc] if loc else []
    L += ['Productgroepen: ' + ', '.join(GROUP_NL.get(g, g) for g in n['groups']), 'Prioriteit: ' + PRIO_NL.get(n['priority'], n['priority'])]
    L += ['Prijs & volume: ' + n['price']] if n['price'] else []
    L += ['', 'Notitie:', n['text']] if n['text'] else []
    if tasks:
        L += ['', 'Acties:'] + ['- [%s] %s (%s%s)' % ('x' if t['done'] else ' ', t['text'], names.get(t['owner_id'], ''),
                                                   ', deadline ' + t['due'] if t['due'] else '') for t in tasks]
    L += (['', 'Bijlagen:'] + ['- ' + x for x in files]) if files else []
    return L + ['']


@app.get('/api/export/zip')
@auth()
def export_zip(u):
    """Map per relatie (visitekaartjes/, fotos/, gescande-notities/, notities.txt) + overzicht.xlsx met per bestand relatie en contact."""
    data = visible_data(u)
    users = {r['id']: r for r in q('SELECT id, name, initials FROM users')}
    names = {k: v['name'] for k, v in users.items()}
    rels = {r['id']: r for r in data['relations']}
    tasks_by_note = {}
    for t in data['tasks']:
        tasks_by_note.setdefault(t['note_id'], []).append(t)
    folders, used, written, index, texts = {}, set(), set(), [], {}

    def folder(rid, name):
        if rid not in folders:
            base = f = safe_name(name)
            i = 2
            while f.lower() in used:
                f, i = '%s (%d)' % (base, i), i + 1
            used.add(f.lower())
            folders[rid] = f
        return folders[rid]

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in sorted(data['notes'], key=lambda x: (x['date'], x['time'])):
            r, fd, files, count = rels.get(n['relation_id'], {}), folder(n['relation_id'], n['company']), [], {}
            for a in n['attachments']:
                fn = a['url'].rsplit('/', 1)[-1]
                path = os.path.join(UPLOADS, fn)
                if not os.path.exists(path):
                    continue
                count[a['kind']] = count.get(a['kind'], 0) + 1
                stem = safe_name('%s %s %s %d' % (n['date'], n['contact'] or 'zonder naam', users.get(n['user_id'], {}).get('initials', ''), count[a['kind']]), 100)
                arc = '%s/%s/%s%s' % (fd, KIND_DIR.get(a['kind'], 'fotos'), stem.replace(' ', '_'), os.path.splitext(fn)[1])
                root, ext = os.path.splitext(arc)
                k = 2
                while arc.lower() in written:
                    arc, k = '%s_%d%s' % (root, k, ext), k + 1
                written.add(arc.lower())
                z.write(path, arc)
                files.append(arc.split('/', 1)[1])
                index.append({'relation': n['company'], 'type': type_nl(n['type']), 'country': n['country'], 'csb': r.get('csb', ''),
                              'contact': n['contact'], 'role': n['contact_role'], 'email': n['contact_email'], 'phone': n['contact_phone'],
                              'kind': KIND_NL.get(a['kind'], a['kind']), 'date': n['date'], 'time': n['time'],
                              'user': names.get(n['user_id'], ''), 'file': arc})
            if fd not in texts:
                head = [n['company'], 'Type: ' + type_nl(n['type'])]
                head += ['Land: ' + n['country']] if n['country'] else []
                head += ['CSB-nummer: ' + r['csb']] if r.get('csb') else []
                head += ['Eigenaar: ' + names.get(r['owner_id'], '')] if r.get('owner_id') else []
                head += ['Fase: ' + STAGE_NL.get(r['stage'], r['stage'])] if r.get('stage') else []
                texts[fd] = head + ['']
            texts[fd] += note_txt(n, names, tasks_by_note.get(n['id'], []), files)
        for fd, lines in texts.items():
            z.writestr(fd + '/notities.txt', '\r\n'.join(lines).encode('utf-8-sig'))
        index.sort(key=lambda x: (x['relation'].lower(), x['date'], x['time']))
        z.writestr('overzicht.xlsx', to_xlsx({'Bijlagen': (IDX_COLS, index)}, link_col='file'))
        z.writestr('overzicht.csv', to_csv(IDX_COLS, index))
        z.writestr('data.json', json.dumps({'exported_at': now(), 'settings': settings(), **data}, ensure_ascii=False, indent=2))
        for title, (cols, items) in export_tables(u).items():
            z.writestr('tabellen/%s.csv' % title.lower(), to_csv(cols, items))
    return download(buf.getvalue(), 'application/zip', 'luiten-crm-%s.zip' % stamp())


# ---------- start ----------
def run_server():
    """Threaded HTTP-server uit de standaardbibliotheek (wsgiref); bedoeld achter Caddy."""
    from socketserver import ThreadingMixIn
    from wsgiref.simple_server import WSGIRequestHandler, WSGIServer, make_server

    class Server(ThreadingMixIn, WSGIServer):
        daemon_threads = True
        allow_reuse_address = True
        request_queue_size = 128  # standaard 5: geeft vertraging bij veel gelijktijdige verzoeken

    class Handler(WSGIRequestHandler):
        timeout = 60  # hangende verbindingen niet eindeloos vasthouden
        server_version, sys_version = 'LuitenCRM', ''

        def log_request(self, code='-', size='-'):  # toegangslog staat al in Caddy; alleen 5xx loggen
            if str(getattr(code, 'value', code)).startswith('5'):
                super().log_request(code, size)

    httpd = make_server(HOST, PORT, app, server_class=Server, handler_class=Handler)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        httpd.server_close()


if __name__ == '__main__':
    if PUBLIC_URL and not PUBLIC:
        sys.exit('LCRM_PUBLIC_URL moet met https:// beginnen.')
    first = init_db()
    line = '=' * 62
    print('\n' + line + '\n  LUITEN CRM %s draait' % APP_VERSION)
    if PUBLIC:
        print('  Publiek adres:  %s   (via HTTPS-proxy, intern op %s:%d)' % (PUBLIC_URL, HOST, PORT))
    else:
        print('  Lokaal (alleen testen): http://localhost:%d  /  %s' % (PORT, app_url()))
    if first:
        print('\n  EERSTE KEER - log in als beheerder:\n    gebruikersnaam: beheerder\n    wachtwoord:     %s' % first)
        print('  (staat ook in %s; je kiest direct een eigen wachtwoord)' % os.path.join(DATA, 'EERSTE-WACHTWOORD.txt'))
    print(line + '\n', flush=True)
    run_server()
